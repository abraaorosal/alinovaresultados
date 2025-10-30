#!/usr/bin/env python3
from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
WORKBOOK_PATH = ROOT / "Planilha de desperdício.xlsx"
OUTPUT_PATH = ROOT / "src/data/waste-data.json"
PRODUCT_SHEET = "CADASTRO PRODUTOS"
MONTH_SHEET_PREFIX = "DESPERDÍCIO"


def normalize_number(value: Optional[float]) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raise TypeError(f"Unsupported numeric value: {value!r}")


@dataclass
class WasteRecord:
    date: datetime
    product: str
    cost_per_kg: Optional[float]
    waste_kg: Optional[float]
    waste_cost: Optional[float]
    reason: Optional[str]

    def as_dict(self) -> dict:
        return {
            "date": self.date.isoformat(),
            "product": self.product,
            "costPerKg": normalize_number(self.cost_per_kg),
            "wasteKg": normalize_number(self.waste_kg),
            "wasteCost": normalize_number(self.waste_cost),
            "reason": self.reason if self.reason else None,
        }


@dataclass
class MonthSheet:
    name: str
    month: datetime
    client: str
    cmv_atual: float
    faturamento_medio: float
    records: List[WasteRecord]

    @property
    def custo_desperdicio(self) -> float:
        return sum(
            record.waste_cost or 0.0
            for record in self.records
        )

    @property
    def cmv_em_reais(self) -> float:
        return self.faturamento_medio * self.cmv_atual

    @property
    def percentual_desperdicio(self) -> float:
        cmv = self.cmv_em_reais
        return (self.custo_desperdicio / cmv) if cmv else 0.0


def load_products(sheet) -> List[dict]:
    header_row_idx = None
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row and row[0] == "PRODUTO":
            header_row_idx = idx
            break
    if header_row_idx is None:
        raise RuntimeError("Cabeçalho 'PRODUTO' não encontrado na aba de produtos.")

    products: List[dict] = []
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        if not row or not row[0]:
            continue
        name = row[0]
        price_per_kg = normalize_number(row[1])
        products.append({"name": name, "pricePerKg": price_per_kg})
    return products


def load_month_sheet(sheet) -> MonthSheet:
    client = ""
    cmv_atual = 0.0
    faturamento_medio = 0.0
    month_date: Optional[datetime] = None

    for row in sheet.iter_rows(min_row=1, max_row=10, values_only=True):
        if not row:
            continue
        label = row[1]
        if label == "CLIENTE:":
            client = row[2]
            cmv_atual = float(row[6] or 0.0)
        elif label == "MÊS:":
            month_date = row[2]
            faturamento_medio = float(row[6] or 0.0)

    if not month_date:
        raise RuntimeError(f"Data do mês não encontrada na aba '{sheet.title}'.")

    header_row_idx = None
    for idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row and "DATA" in row:
            header_row_idx = idx
            break
    if header_row_idx is None:
        raise RuntimeError(f"Cabeçalho de dados não encontrado na aba '{sheet.title}'.")

    DATA_COL = 1
    PRODUTO_COL = 2
    CUSTO_KG_COL = 4
    PESO_COL = 5
    CUSTO_COL = 6
    MOTIVO_COL = 8

    records: List[WasteRecord] = []
    for row in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        date_cell = row[DATA_COL] if len(row) > DATA_COL else None
        product_cell = row[PRODUTO_COL] if len(row) > PRODUTO_COL else None

        if not date_cell or not product_cell:
            continue

        if not isinstance(date_cell, datetime):
            raise ValueError(f"Data inválida na aba '{sheet.title}': {date_cell!r}")

        records.append(
            WasteRecord(
                date=date_cell,
                product=str(product_cell),
                cost_per_kg=normalize_number(row[CUSTO_KG_COL]) if len(row) > CUSTO_KG_COL else None,
                waste_kg=normalize_number(row[PESO_COL]) if len(row) > PESO_COL else None,
                waste_cost=normalize_number(row[CUSTO_COL]) if len(row) > CUSTO_COL else None,
                reason=str(row[MOTIVO_COL]).strip() if len(row) > MOTIVO_COL and row[MOTIVO_COL] else None,
            )
        )

    return MonthSheet(
        name=sheet.title,
        month=month_date,
        client=client,
        cmv_atual=cmv_atual,
        faturamento_medio=faturamento_medio,
        records=records,
    )


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {WORKBOOK_PATH}")

    workbook = load_workbook(WORKBOOK_PATH, data_only=True)

    product_sheet = workbook[PRODUCT_SHEET]
    products = load_products(product_sheet)

    month_sheets: List[MonthSheet] = []
    for sheet_name in workbook.sheetnames:
        if sheet_name.startswith(MONTH_SHEET_PREFIX):
            month_sheets.append(load_month_sheet(workbook[sheet_name]))

    if not month_sheets:
        raise RuntimeError("Nenhuma aba de desperdício encontrada na planilha.")

    month_sheets.sort(key=lambda item: item.month)
    latest = month_sheets[-1]

    waste_records = sorted(
        (record for month in month_sheets for record in month.records),
        key=lambda record: record.date,
    )

    dataset = {
        "clientInfo": {
            "client": latest.client,
            "cmvAtual": latest.cmv_atual,
            "month": latest.month.isoformat(),
            "faturamentoMedio": latest.faturamento_medio,
            "cmvEmReais": latest.cmv_em_reais,
            "custoDesperdicio": latest.custo_desperdicio,
            "percentualDesperdicio": latest.percentual_desperdicio,
        },
        "products": products,
        "wasteRecords": [record.as_dict() for record in waste_records],
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(dataset, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Dataset atualizado com {len(waste_records)} registros em {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
